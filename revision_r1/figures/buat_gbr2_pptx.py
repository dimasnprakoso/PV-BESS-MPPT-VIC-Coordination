# -*- coding: utf-8 -*-
"""
Gambar 2 (sistem studi) sebagai bentuk PowerPoint ASLI yang dapat disunting (komentar R1-6:
"Figure 1 is a block diagram and should be supplied in an editable format").
Tata letak menyalin gbr1_sistem (sumber TikZ fig1_arch.tex); koordinat dalam mm pada kanvas 200 x 101 mm.
Tidak memakai PowerPoint COM (tidak menyentuh PowerPoint yang sedang dibuka pengguna).
Keluaran: gbr2_sistem.pptx + gbr2_sistem_pratinjau.png (pratinjau matplotlib dari daftar elemen yang sama).
Teks: X_{sub} = subskrip, *...* = miring.
"""
import os
import re

from pptx import Presentation
from pptx.util import Mm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
W_MM, H_MM = 200, 101
HITAM, PUTIH, ABU = RGBColor(0, 0, 0), RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xF2, 0xF2, 0xF2)
FONT = "Times New Roman"
TEBAL, TIPIS = 2.6, 0.9          # pt: jalur daya, jalur sinyal

# ---------------------------------------------------------------- daftar elemen (mm)
# kotak: (x1, y1, x2, y2, teks, abu?, ukuran)
KOTAK = [
    (9.0, 7.8, 30.3, 21.3, "PV\narray", False, 10),
    (41.7, 7.8, 60.2, 21.3, "", False, 10),             # boost (diagonal + DC/DC)
    (104.3, 7.8, 122.7, 21.3, "", False, 10),           # inverter (diagonal + DC/AC)
    (132.7, 7.8, 154.0, 21.3, "LCL\nfilter", False, 10),
    (177.5, 7.8, 198.7, 21.3, "Grid\n50 Hz", False, 10),
    (44.2, 32.3, 57.8, 39.5, "PWM", False, 9),
    (68.7, 28.5, 95.8, 42.0, "Bidirectional\nDC–DC", False, 10),
    (106.7, 28.1, 120.3, 35.2, "PWM", False, 9),
    (55.6, 46.6, 69.1, 53.7, "PWM", False, 9),
    (71.6, 50.5, 92.9, 64.0, "BESS", False, 10),
    (9.0, 50.8, 30.3, 63.6, "MPPT", True, 10),
    (101.3, 50.8, 125.7, 63.6, "*dq*-frame\ncurrent ctrl", True, 10),
    (131.6, 50.8, 149.4, 63.6, "VIC", True, 10),
    (159.4, 52.8, 167.3, 61.6, "*d*/*dt*", False, 8),
    (177.5, 50.8, 198.7, 63.6, "PLL", True, 10),
    (59.8, 75.0, 104.6, 87.8, "Adaptive coordination\n*α*, *β*", True, 10),
]
# garis: (titik-titik, tebal?, putus?, panah: 0 tanpa / 1 ujung / 2 dua arah)
GARIS = [
    ([(30.3, 14.7), (41.7, 14.7)], True, False, 1),
    ([(60.2, 14.7), (104.3, 14.7)], True, False, 1),
    ([(122.7, 14.7), (132.7, 14.7)], True, False, 1),
    ([(154.0, 14.7), (177.5, 14.7)], True, False, 1),
    ([(82.3, 14.7), (82.3, 28.5)], True, False, 2),
    ([(82.3, 42.0), (82.3, 50.5)], True, False, 2),
    ([(69.4, 14.7), (69.4, 19.0)], False, False, 0),           # kapasitor C_dc
    ([(66.3, 19.0), (72.5, 19.0)], False, False, 0),
    ([(66.3, 20.8), (72.5, 20.8)], False, False, 0),
    ([(51.0, 32.3), (51.0, 21.3)], False, False, 1),           # PWM -> boost
    ([(19.7, 50.8), (19.7, 48.0), (51.0, 48.0), (51.0, 39.5)], False, False, 1),   # MPPT -> PWM
    ([(113.5, 28.1), (113.5, 21.3)], False, False, 1),         # PWM -> inverter
    ([(113.5, 50.8), (113.5, 35.2)], False, False, 1),         # dq -> PWM
    ([(62.3, 46.6), (62.3, 35.2), (68.7, 35.2)], False, False, 1),   # PWM -> bidirectional
    ([(62.3, 75.0), (62.3, 53.7)], False, False, 1),           # koordinasi -> PWM (beta P_bat)
    ([(82.3, 64.0), (82.3, 75.0)], False, True, 1),            # SOC (putus-putus)
    ([(19.7, 63.6), (19.7, 81.5), (59.8, 81.5)], False, False, 1),   # P_mppt
    ([(104.6, 77.0), (113.5, 77.0), (113.5, 63.6)], False, False, 1),  # P_ref
    ([(140.5, 63.6), (140.5, 85.7), (104.6, 85.7)], False, False, 1),  # P_vic
    ([(159.4, 57.2), (149.4, 57.2)], False, False, 1),         # d/dt -> VIC
    ([(177.5, 57.2), (167.3, 57.2)], False, False, 1),         # PLL -> d/dt
    ([(173.2, 57.2), (173.2, 96.3), (82.3, 96.3), (82.3, 87.8)], False, False, 1),   # delta omega
    ([(188.3, 63.6), (188.3, 99.3), (117.8, 99.3), (117.8, 63.6)], False, False, 1), # theta
    ([(9.0, 14.7), (1.2, 14.7), (1.2, 57.2), (9.0, 57.2)], False, True, 1),         # V_pv, I_pv
    ([(165.5, 14.7), (165.5, 31.6)], False, True, 0),          # pengukuran PCC
    ([(165.5, 31.6), (125.0, 31.6), (125.0, 50.8)], False, True, 1),
    ([(165.5, 31.6), (188.3, 31.6), (188.3, 50.8)], False, True, 1),
]
TITIK = [(69.4, 14.7), (82.3, 14.7), (165.5, 14.7), (165.5, 31.6), (173.2, 57.2)]
# label: (x tengah, y tengah, teks, ukuran)
LABEL = [
    (51.0, 4.0, "Boost converter", 9), (113.5, 4.0, "Inverter", 9),
    (88.5, 11.0, "*V*_{dc}", 9), (76.0, 23.5, "*C*_{dc}", 9), (165.5, 11.2, "PCC", 9),
    (35.5, 44.5, "*duty*", 9), (120.0, 24.5, "*g*_{pwm}", 9),
    (141.0, 28.8, "*v*_{abc}, *i*_{abc}", 9), (178.5, 28.8, "*v*_{abc}", 9),
    (9.7, 35.7, "*V*_{pv}, *I*_{pv}", 9), (88.5, 69.5, "*SOC*", 9),
    (56.0, 64.8, "*βP*_{bat}", 9), (39.7, 77.8, "*P*_{mppt}", 9), (108.0, 73.8, "*P*_{ref}", 9),
    (146.0, 74.8, "*P*_{vic}", 9), (154.5, 53.2, "Δ*ω̇*", 9), (172.3, 53.2, "Δ*ω*", 9),
    (167.7, 93.5, "Δ*ω*", 9), (183.8, 96.5, "*θ*", 9),
    (128.0, 14.0, "///", 8), (158.5, 14.0, "///", 8),
]
KOTAK_DIAG = [((41.7, 7.8, 60.2, 21.3), "DC", "DC"), ((104.3, 7.8, 122.7, 21.3), "DC", "AC")]


def potong(teks):
    """'*V*_{dc}' -> [(teks, miring, subskrip), ...]"""
    out = []
    for bag in re.split(r"(_\{[^}]*\})", teks):
        if not bag:
            continue
        sub = bag.startswith("_{")
        if sub:
            bag = bag[2:-1]
        for k, s in enumerate(re.split(r"\*", bag)):
            if s:
                out.append((s, k % 2 == 1 or (sub and out and out[-1][1]), sub))
    return out


# ---------------------------------------------------------------- PowerPoint
prs = Presentation()
prs.slide_width, prs.slide_height = Mm(W_MM), Mm(H_MM)
sl = prs.slides.add_slide(prs.slide_layouts[6])


def isi_teks(tf, teks, size):
    tf.word_wrap = False
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    for k, baris in enumerate(teks.split("\n")):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        for s, it, sub in potong(baris):
            r = p.add_run(); r.text = s
            r.font.size = Pt(size); r.font.name = FONT; r.font.italic = it; r.font.color.rgb = HITAM
            if sub:
                r._r.get_or_add_rPr().set("baseline", "-25000")


def kotak(x1, y1, x2, y2, teks, abu, size):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Mm(x1), Mm(y1), Mm(x2 - x1), Mm(y2 - y1))
    s.fill.solid(); s.fill.fore_color.rgb = ABU if abu else PUTIH
    s.line.color.rgb = HITAM; s.line.width = Pt(TIPIS); s.shadow.inherit = False
    s.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    if teks:
        isi_teks(s.text_frame, teks, size)


def label(cx, cy, teks, size, w=30.0, h=5.0):
    tb = sl.shapes.add_textbox(Mm(cx - w / 2), Mm(cy - h / 2), Mm(w), Mm(h))
    tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    isi_teks(tb.text_frame, teks, size)


def garis(pts, tebal, putus, panah):
    for i in range(len(pts) - 1):
        (x1, y1), (x2, y2) = pts[i], pts[i + 1]
        c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Mm(x1), Mm(y1), Mm(x2), Mm(y2))
        c.line.color.rgb = HITAM; c.line.width = Pt(TEBAL if tebal else TIPIS); c.shadow.inherit = False
        if putus:
            c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        ln = c.line._get_or_add_ln()
        if panah and i == len(pts) - 2:
            ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
        if panah == 2 and i == 0:
            ln.append(ln.makeelement(qn("a:headEnd"), {"type": "triangle", "w": "med", "len": "med"}))


for k in KOTAK:
    kotak(*k)
for (x1, y1, x2, y2), a, b in KOTAK_DIAG:
    garis([(x1, y1), (x2, y2)], False, False, 0)
    label(x1 + 3.2, y1 + 1.7, a, 6, w=6, h=3); label(x2 - 3.2, y2 - 1.7, b, 6, w=6, h=3)
for g in GARIS:
    garis(*g)
for x, y in TITIK:
    d = sl.shapes.add_shape(MSO_SHAPE.OVAL, Mm(x - 0.7), Mm(y - 0.7), Mm(1.4), Mm(1.4))
    d.fill.solid(); d.fill.fore_color.rgb = HITAM; d.line.fill.background(); d.shadow.inherit = False
for lb in LABEL:
    label(*lb)
prs.save(os.path.join(HERE, "gbr2_sistem.pptx"))

# ---------------------------------------------------------------- pratinjau (matplotlib, elemen yang sama)
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle, FancyArrowPatch

plt.rcParams.update({"font.family": "serif", "font.serif": ["Times New Roman"], "mathtext.fontset": "stix"})
fig, ax = plt.subplots(figsize=(W_MM / 25.4, H_MM / 25.4))
ax.set_xlim(0, W_MM); ax.set_ylim(H_MM, 0); ax.axis("off"); fig.subplots_adjust(0, 0, 1, 1)


def mt(teks):
    """ke mathtext untuk pratinjau"""
    out = ""
    for s, it, sub in potong(teks):
        s2 = s.replace(" ", r"\ ")
        body = (s2 if it else r"\mathrm{" + s2 + "}")
        out += ("_{" + body + "}") if sub else body
    return "$" + out + "$" if any(c in teks for c in "*_") else teks


for x1, y1, x2, y2, t, abu, size in KOTAK:
    ax.add_patch(Rectangle((x1, y1), x2 - x1, y2 - y1, fc="#f2f2f2" if abu else "w", ec="k", lw=TIPIS))
    if t:
        ax.text((x1 + x2) / 2, (y1 + y2) / 2, "\n".join(mt(b) for b in t.split("\n")), ha="center", va="center", fontsize=size)
for (x1, y1, x2, y2), a, b in KOTAK_DIAG:
    ax.plot([x1, x2], [y1, y2], "k", lw=TIPIS)
    ax.text(x1 + 3.2, y1 + 1.7, a, ha="center", va="center", fontsize=6); ax.text(x2 - 3.2, y2 - 1.7, b, ha="center", va="center", fontsize=6)
for pts, tebal, putus, panah in GARIS:
    lw = TEBAL if tebal else TIPIS; ls = "--" if putus else "-"
    for i in range(len(pts) - 1):
        last = i == len(pts) - 2
        style = "<|-|>" if (panah == 2) else ("-|>" if (panah and last) else "-")
        ax.add_patch(FancyArrowPatch(pts[i], pts[i + 1], arrowstyle=style, mutation_scale=7, lw=lw, ls=ls, color="k", shrinkA=0, shrinkB=0))
for x, y in TITIK:
    ax.plot(x, y, "ko", ms=2.5)
for cx, cy, t, size in LABEL:
    ax.text(cx, cy, mt(t), ha="center", va="center", fontsize=size)
fig.savefig(os.path.join(HERE, "gbr2_sistem_pratinjau.png"), dpi=300)
print("ok")
